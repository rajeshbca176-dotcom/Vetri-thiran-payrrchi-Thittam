from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from utils import sanitize_text, split_terms

ROOT = Path(__file__).resolve().parents[1]
LOGO = ROOT / "assets" / "legalease_logo.svg"


def _set_cell_shading(cell, fill="1F2937"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def format_docx(text: str, doc_type: str, terms: str = "") -> bytes:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.7)
    sec.bottom_margin = Inches(0.7)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)

    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(11)
    styles["Title"].font.name = "Times New Roman"
    styles["Title"].font.size = Pt(20)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # SVG cannot be embedded by python-docx directly; use a text brand header.
    r = p.add_run("LEGAL EASE")
    r.bold = True
    r.font.size = Pt(16)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(sanitize_text(doc_type).upper())
    run.bold = True
    run.font.size = Pt(18)

    for raw in sanitize_text(text).splitlines():
        line = raw.strip()
        if not line:
            doc.add_paragraph()
            continue
        p = doc.add_paragraph()
        if line.startswith("#"):
            line = line.lstrip("# ")
            p.style = doc.styles["Heading 2"]
            p.add_run(line)
        elif line.isupper() or line.startswith("ARTICLE ") or line.startswith("SECTION "):
            p.style = doc.styles["Heading 2"]
            p.add_run(line)
        else:
            p.add_run(line)

    if terms.strip():
        doc.add_paragraph("Key Terms", style="Heading 2")
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "No."
        table.rows[0].cells[1].text = "Term"
        for cell in table.rows[0].cells:
            _set_cell_shading(cell)
            for run in cell.paragraphs[0].runs:
                run.font.bold = True
                run.font.color.rgb = __import__("docx").shared.RGBColor(255, 255, 255)
        for i, term in enumerate(split_terms(terms), 1):
            cells = table.add_row().cells
            cells[0].text = str(i)
            cells[1].text = term

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("LegalEase — AI-generated draft. Review with a qualified legal professional before use.").italic = True

    out = BytesIO()
    doc.save(out)
    return out.getvalue()
